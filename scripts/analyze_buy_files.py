import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    text = str(value).strip()
    return text[:160] if text else None


def score_row(values):
    tokens = ("po", "style", "sku", "material", "color", "colour", "size", "qty", "quantity", "delivery", "factory", "vendor")
    texts = [str(v).strip().lower() for v in values if v not in (None, "")]
    return len(texts) + 3 * sum(any(token in text for token in tokens) for text in texts)


def analyze(path):
    wb = load_workbook(path, read_only=False, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        max_scan_row = min(ws.max_row or 1, 60)
        max_scan_col = min(ws.max_column or 1, 80)
        candidates = []
        for row_idx in range(1, max_scan_row + 1):
            values = [ws.cell(row_idx, col_idx).value for col_idx in range(1, max_scan_col + 1)]
            candidates.append((score_row(values), row_idx, values))
        candidates.sort(reverse=True, key=lambda item: item[0])
        _, header_row, header_values = candidates[0]
        nonempty_header = [(idx + 1, clean(v)) for idx, v in enumerate(header_values) if clean(v) is not None]
        samples = []
        for row_idx in range(header_row + 1, min(ws.max_row, header_row + 4) + 1):
            row = [clean(ws.cell(row_idx, col_idx).value) for col_idx in range(1, max_scan_col + 1)]
            if any(v is not None for v in row):
                samples.append(row)
        sheets.append({
            "name": ws.title,
            "rows": ws.max_row,
            "columns": ws.max_column,
            "merged_ranges": [str(rng) for rng in list(ws.merged_cells.ranges)[:20]],
            "likely_header_row": header_row,
            "headers": nonempty_header,
            "sample_rows": samples,
            "top_candidates": [{"row": row, "score": score} for score, row, _ in candidates[:5]],
        })
    return {"file": str(path), "sheets": sheets}


if __name__ == "__main__":
    compact = "--compact" in sys.argv
    results = []
    for arg in [item for item in sys.argv[1:] if item != "--compact"]:
        try:
            result = analyze(Path(arg))
            if compact:
                result = {
                    "file": Path(result["file"]).name,
                    "sheets": [
                        {
                            "name": sheet["name"],
                            "rows": sheet["rows"],
                            "columns": sheet["columns"],
                            "likely_header_row": sheet["likely_header_row"],
                            "headers": [value for _, value in sheet["headers"]],
                        }
                        for sheet in result["sheets"]
                    ],
                }
            results.append(result)
        except Exception as exc:
            results.append({"file": arg, "error": f"{type(exc).__name__}: {exc}"})
    print(json.dumps(results, ensure_ascii=False, indent=2))
